from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
st = d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10.5)
TITLE = "Heirline: Digital Inheritance Claims Idea (Research, Oct 2026)"
d.add_heading(TITLE, 0)
d.add_paragraph("Working doc for Austin Thomas. Research notes on a side-business idea. Facts only; sources linked inline.")

def h(t): d.add_heading(t, 1)
def p(t): d.add_paragraph(t)
def b(t, bold=None):
    para = d.add_paragraph(style='List Bullet')
    if bold:
        r = para.add_run(bold); r.bold = True
    para.add_run(t)
def shade(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),'D9E2F3'); tcPr.append(s)
def table(headers, rows):
    t = d.add_table(rows=1, cols=len(headers)); t.style='Table Grid'
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(x); r.bold=True; shade(c)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):
            cells[i].text=x
            for para in cells[i].paragraphs:
                for r in para.runs: r.font.size=Pt(9)
    d.add_paragraph()

h("1. Idea summary")
p("B2B software for 401(k) recordkeepers, IRA/HSA custodians, and life insurers. Customers pay per covered member per year.")
b(" lists the customer's accounts, finds missing or outdated beneficiaries, and backs up crypto keys through partners (Casa, Vault12).", "While the customer is alive:")
b(" beneficiaries get a guided claim, a filled-in claim packet for each custodian, status tracking, and plain-English guidance on payout choices.", "After a death:")
b(" families can already get a 'one-stop shop' from Empathy, so the defensible piece is the layer that actually files claims with each custodian.", "Strategic point:")

h("2. Market")
b(" $84.4T in wealth transfers through 2045, $72.6T to heirs. Source: https://www.cerulli.com/press-releases/cerulli-anticipates-84-trillion-in-wealth-transfers-through-2045", "Cerulli projection:")
b(" 31.9M forgotten 401(k)s holding $2.1T (July 2025).", "Capitalize's own analysis:")
b(" 60M+ US deaths expected over 20 years.", "Census data cited by 2048 Ventures:")

h("3. Competitors")
table(["Name","What it does","Price","Funding","URL"],[
["Empathy","After-death app plus planning vault, sold through employers, insurers and banks; says it covers 40M+ people. Partnered with Voya. Insurer backers include Aflac, MetLife, New York Life, TIAA.","Premium $8.99/mo or $64.99/yr","$160M+ incl. $72M Series C (May 2025)","https://www.empathy.com/about"],
["Elayne","Says it discovers accounts, files claims, and manages distributions. Closest competitor.","Unknown","Unknown","https://www.elayne.com/solutions/families"],
["Settld","UK; sends death notices and closes accounts.","—","—","https://www.settld.care/"],
["Trustworthy","Household vault.","Free, then $10/$20/$40 per month billed annually","$19.7M","https://www.trustworthy.com/pricing"],
["Everplans","Shareable vault.","$99.99/yr","—","https://www.everplans.com/pricing"],
["Lantern","Planning and after-death checklists.","—","$2.3M","https://www.lantern.co"],
["Atticus","Model, not competitor: a law firm that wins benefit claims for a contingency fee.","Contingency fee","$100M+","https://www.atticus.com"],
["Casa","Multisig crypto inheritance.","$250 to $7,500/yr","—","https://casa.io/inheritance"],
["Vault12","Seed phrase split among people you choose.","$29.99/mo inheritance plan","—","https://vault12.com/pricing/"],
["Ledger Recover","Seed backup with Coincover.","—","—","https://store.ledger.com/products/ledger-recover"],
["Capitalize","Finds old 401(k)s and rolls them into IRAs.","—","$19M Series B (Aug 2024)","—"],
])

h("4. Custodian claim processes")
table(["Custodian","How to start","Documents","Third-party rules","Timeline","URL"],[
["Schwab","Report the death online; Inheritance Center with to-do lists, document uploads and status.","Death certificate; may need a notarized affidavit of domicile or a state tax waiver.","—","Death certificate verified in about 5 business days; most transfers take a few weeks.","https://www.schwab.com/life-events/losing-a-loved-one"],
["Fidelity","Report online with the deceased's SSN; claim form sent by mail or uploaded in the NetBenefits app.","Claim form; may need a medallion signature guarantee (a notary doesn't count).","A POA holder or guardian can sign with proof; some 401(k)s need plan-sponsor approval.","—","https://digital.fidelity.com/prgw/digital/journeys/loss/report-a-death"],
["Vanguard","Online for some accounts if you're the named beneficiary or executor; otherwise phone 877-662-7447 or a paper kit.","Certified death certificate; estates need a current court appointment.","—","—","https://investor.vanguard.com/inheriting-accounts"],
["Empower","Notification form plus death certificate, or phone 866-442-3888.","Notification form; death certificate.","Accepts notice from anyone with knowledge of the death.","About one month after documents are complete.","https://www.empower.com/beneficiary-claim"],
["HealthEquity","Instructions Upon Death form plus death certificate.","Instructions Upon Death form; death certificate.","Investments are sold to cash first; a spouse can move the money into their own HSA; anyone else gets a check, possibly minus a closure fee.","Form reviewed within 2 business days.","https://www.healthequity.com/help-center/articles/6026409-hsa-account-management"],
["Coinbase","Filed through the Executor Services form.","Probate papers or a small-estate affidavit; photo ID for each named person; signed letter naming the Coinbase account to move assets to.","Can't name a beneficiary on individual accounts.","—","https://help.coinbase.com/en/coinbase/managing-my-account/other/how-do-i-gain-access-to-a-deceased-family-members-coinbase-account"],
])
para = d.add_paragraph(); r = para.add_run("None of the six publishes an API for death claims."); r.bold = True

h("5. What software can automate vs. what needs humans or partners")
d.add_heading("Software can automate", 2)
for x in ["Filling out forms from one set of details","A document checklist for each custodian, including how many death certificates to order","Online death reports","Tracking and reminders","Drafting the Coinbase transfer letter","Checking beneficiaries while the person is alive and flagging crypto that will go through probate"]: b(x)
d.add_heading("Needs humans or partners", 2)
for x in ["Wet signatures and medallion guarantees","Court letters","Plan-sponsor approvals","Identity checks inside each custodian's system","Acting on someone's behalf (needs a POA or the beneficiary's own login)","Direct data access (needs a partner agreement)"]: b(x)

h("6. Risks")
for x in ["Custodians may refuse to deal with a third party","State 'heir finder' laws on taking a cut of recovered money","Rules on investment advice and on practicing law","GLBA privacy rules for SSNs and financial data","Security, since the service holds a full map of someone's assets","Competition from Empathy and Elayne"]: b(x)

h("7. Lost crypto key note")
b("A lost private key can't be recovered.")
b(" a 2-of-3 multisig wallet (any two of three keys can move funds) or a Shamir seed split (Trezor).", "Prevent it with")
b("Password-guessing recovery services only help if you still have the wallet file.")

h("8. Next steps")
for x in ["Try Elayne and Empathy.","Interview 10 people who handle beneficiary claims at plan sponsors or IRA custodians.","Pitch a free beneficiary audit to one small recordkeeper or credit union.","First partners to target: 401(k) recordkeepers (Voya, Empower), then HSA custodians, with crypto last."]: b(x)
b(" SE experience with enterprise sales and integration, selling to security-conscious buyers.", "Austin's edge:")

d.save('/workspace/heirline/heirline.docx')
